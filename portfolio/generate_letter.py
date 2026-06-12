"""Generate a motivation letter (cover letter) in PDF and Word (.docx).

A strong, ATS-friendly letter built automatically from Gerrit's profile and how
it relates to the vacancy — no manual "why this company" text needed. Available
in Dutch and English. Run:

    python generate_letter.py        # NL + EN templates with [placeholders]

Or import and call build_pdf/build_docx with company/role/contact/highlights
filled in (``highlights`` = the keywords the vacancy and CV share).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

INK = HexColor("#1f2a44")
SENDER = "Gerrit Düthler"
SENDER_CONTACT = "gerrit@duthler.info · Eindhoven · linkedin.com/in/gerrit-d-90a50234"

DEFAULTS = {
    "nl": {"company": "[Bedrijf]", "role": "[Functie]", "contact": "heer/mevrouw"},
    "en": {"company": "[Company]", "role": "[Role]", "contact": "Sir or Madam"},
}


# Proper casing for acronyms/brand terms (applied in both languages).
_CASE = {
    "power bi": "Power BI", "sap": "SAP", "vba": "VBA", "cnc": "CNC", "fmea": "FMEA",
    "dmaic": "DMAIC", "5s": "5S", "six sigma": "Six Sigma", "green belt": "Green Belt",
    "lean": "Lean", "excel": "Excel", "python": "Python", "kaizen": "Kaizen",
}
# Dutch keyword → English term, so the English letter doesn't carry Dutch words.
_EN = {
    "calculatie": "cost estimating", "calculator": "cost estimating",
    "kostprijs": "cost pricing", "kostencalculatie": "cost estimating",
    "nacalculatie": "post-calculation", "offertecalculatie": "quotation costing",
    "werkvoorbereider": "work preparation", "werkvoorbereiding": "work preparation",
    "industrialisatie": "industrialisation", "maakstrategie": "make strategy",
    "procesverbetering": "process improvement", "inkoop": "purchasing",
    "cost engineer": "cost engineering", "cost estimator": "cost estimating",
    "manufacturing engineer": "manufacturing engineering",
}


def _term(keyword: str, lang: str) -> str:
    kw = (keyword or "").strip().lower()
    if lang == "en" and kw in _EN:
        return _EN[kw]
    return _CASE.get(kw, keyword)


def _highlight_phrase(highlights, lang: str) -> str:
    """Turn the shared CV/vacancy keywords into a clean, language-correct phrase."""
    items, seen = [], set()
    for h in highlights or []:
        t = _term(h, lang)
        if t and t.lower() not in seen:
            seen.add(t.lower())
            items.append(t)
    items = items[:5]
    if not items:
        return ("kostencalculatie, should-cost en procesverbetering" if lang == "nl"
                else "cost estimating, should-cost and process improvement")
    if len(items) == 1:
        return items[0]
    joiner = " en " if lang == "nl" else " and "
    return ", ".join(items[:-1]) + joiner + items[-1]


def _default_core(company: str, role: str, hl: str, lang: str) -> list[str]:
    """The fallback body when no AI text is supplied — aligned to the role, no jargon list."""
    if lang == "en":
        return [
            (f"Cost engineering is my craft: I turn engineering into cost prices a business can "
             f"act on. With a career in high-tech manufacturing — DAF Trucks, VDL ETG, Andritz "
             f"and Wärtsilä — I am applying with genuine enthusiasm for the role of {role} at "
             f"{company}."),
            (f"Your vacancy centres on {hl}, and that is exactly where I am at my best. I make "
             f"cost prices defensible from first quote to post-calculation, keep estimating fast "
             f"and transparent, and protect margin — bringing engineering, purchasing, sales and "
             f"business control onto one and the same number. Where it helps, I build the tools "
             f"that sharpen the calculation myself."),
            (f"I would gladly show you in a personal conversation what that looks like in "
             f"practice — and what it could mean for {company}. You can reach me at "
             f"gerrit@duthler.info."),
        ]
    return [
        (f"Cost engineering is mijn vak: ik vertaal techniek naar kostprijzen waar een bedrijf "
         f"op kan sturen. Met een loopbaan in de high-tech maakindustrie — DAF Trucks, VDL ETG, "
         f"Andritz en Wärtsilä — solliciteer ik met veel enthousiasme naar de functie van {role} "
         f"bij {company}."),
        (f"Uw vacature draait om {hl}, en daar ben ik op mijn best. Ik maak kostprijzen "
         f"onderbouwd van eerste offerte tot nacalculatie, houd calculeren snel en transparant, "
         f"en bewaak de marge — met engineering, inkoop, verkoop en business control op één en "
         f"hetzelfde cijfer. Waar het helpt, bouw ik zelf de tools die de calculatie scherper "
         f"maken."),
        (f"In een persoonlijk gesprek laat ik u graag zien hoe dat er in de praktijk uitziet — "
         f"en wat het voor {company} kan betekenen. U kunt mij bereiken via "
         f"gerrit@duthler.info."),
    ]


def paragraphs(company: str, role: str, contact: str, highlights=None, lang: str = "nl",
               body=None) -> list[str]:
    """Full letter = letterhead + salutation + core message + sign-off.

    ``body`` (a list of paragraph strings) lets a smarter, vacancy-aligned text be slotted
    in; when omitted, a clean default that matches the role is used.
    """
    today = date.today().strftime("%d-%m-%Y")
    hl = _highlight_phrase(highlights, lang)
    core = [p for p in (body or []) if str(p).strip()] or _default_core(company, role, hl, lang)
    if lang == "en":
        head = [f"Eindhoven, {today}", f"{company}", f"Attn: {contact}",
                f"Subject: application for {role}", f"Dear {contact},"]
        foot = ["Kind regards,", SENDER, SENDER_CONTACT]
    else:
        head = [f"Eindhoven, {today}", f"{company}", f"T.a.v. {contact}",
                f"Betreft: sollicitatie {role}", f"Geachte {contact},"]
        foot = ["Met vriendelijke groet,", SENDER, SENDER_CONTACT]
    return head + core + foot


def _resolve(company, role, contact, lang):
    d = DEFAULTS["en" if lang == "en" else "nl"]
    return (company or d["company"], role or d["role"], contact or d["contact"])


def build_pdf(path: Path, company="", role="", contact="", highlights=None, lang="nl", body=None) -> None:
    company, role, contact = _resolve(company, role, contact, lang)
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("body", parent=styles["Normal"], fontSize=10.5, leading=15,
                                textColor=INK, alignment=TA_LEFT, spaceAfter=8)
    paras = paragraphs(company, role, contact, highlights, lang, body)
    flow = []
    for i, p in enumerate(paras):
        flow.append(Paragraph(p, body_style))
        if i in (0, 2, 3):           # extra space after date block, address, subject
            flow.append(Spacer(1, 6))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm,
                      topMargin=20 * mm, bottomMargin=18 * mm,
                      title=f"Motivatiebrief {SENDER}", author=SENDER).build(flow)


def build_docx(path: Path, company="", role="", contact="", highlights=None, lang="nl", body=None) -> None:
    from docx import Document
    from docx.shared import Pt

    company, role, contact = _resolve(company, role, contact, lang)
    doc = Document()
    f = doc.styles["Normal"].font
    f.name, f.size = "Calibri", Pt(11)
    for p in paragraphs(company, role, contact, highlights, lang, body):
        doc.add_paragraph(p)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


if __name__ == "__main__":
    assets = Path(__file__).parent / "assets"
    build_pdf(assets / "motivation.pdf", lang="nl")
    build_docx(assets / "motivation.docx", lang="nl")
    build_pdf(assets / "motivation_en.pdf", lang="en")
    build_docx(assets / "motivation_en.docx", lang="en")
    print(f"Wrote NL + EN motivation letters in {assets}")
